"""Tests for native-text (PDF text-layer / DOCX) field extraction."""
from __future__ import annotations

from app.extractors.base import RawDocument, RawPage
from app.extractors.text_extractor import TextExtractor, extract_fields_from_text
from app.models import DocumentType

_BANK_EN = """Bank Statement
Statement No.
BS-2026-001
Account Holder
Gulf Foods Import Co.
Currency
USD
Opening Balance 125,000.00
Closing Balance 142,450.00
"""

_COO_AR = """شهادة منشأ
رقم الشهادة  |  شهادة-منشأ-001
المُصدّر  |  شركة كامارو للتجارة ذ.م.م
بلد المنشأ  |  الهند
تاريخ الإصدار  |  2026-10-06
"""


def test_english_bank_statement_fields():
    f = extract_fields_from_text(_BANK_EN, DocumentType.BANK_STATEMENT)
    assert f["opening_balance"].value == "125,000.00"
    assert f["closing_balance"].value == "142,450.00"
    assert f["currency"].value == "USD"
    assert f["statement_no"].value == "BS-2026-001"
    assert f["account_holder"].value == "Gulf Foods Import Co."
    # No mock/placeholder values leaked in.
    assert all("ACME-INV" not in str(fld.value) for fld in f.values())


def test_arabic_keys_and_values_preserved():
    f = extract_fields_from_text(_COO_AR, DocumentType.CERTIFICATE_OF_ORIGIN)
    # Arabic labels map to the canonical names the rules check; the label is kept for display.
    assert f["certificate_number"].value == "شهادة-منشأ-001"
    assert f["certificate_number"].label == "رقم الشهادة"
    assert f["country_of_origin"].value == "الهند"
    assert f["issue_date"].value == "2026-10-06"
    assert f["exporter"].value == "شركة كامارو للتجارة ذ.م.م"


def test_text_extractor_populates_full_text_and_engine():
    raw = RawDocument(name="stmt.docx", pages=[RawPage(0, _BANK_EN)], text_is_native=True)
    ex = TextExtractor().extract(raw, DocumentType.BANK_STATEMENT)
    assert ex.engine.startswith("text:")
    assert len(ex.full_text) >= 5
    assert ex.fields  # real fields, not empty


def test_default_extractor_never_serves_mock_for_real_uploads(monkeypatch):
    # Deployments ship no .env, so the code default must extract real content.
    from app.config import Settings
    from app.extractors import base

    monkeypatch.setattr(base, "settings", Settings(_env_file=None))
    raw = RawDocument(name="stmt.docx", pages=[RawPage(0, _BANK_EN)], text_is_native=True)
    assert base.get_extractor(raw).name == "text"


# --- QA sample-document regressions -------------------------------------------
_PO_EN = """Purchase Order
SAMPLE / DEMO DOCUMENT — For testing and UAT only
PO No.  |  PO-2026-5001
Supplier  |  Camaroo Trading LLC
Order Value  |  USD 24,850.00
"""

_INVOICE_EN = """Commercial Invoice
Invoice No.  |  CI-2026-1001
Invoice Date  |  2026-10-06
Currency  |  USD
Total Amount  |  24,850.00
Description  |  HS Code  |  Quantity  |  Unit Price  |  Amount
Frozen Shrimp  |  030617  |  500 Cartons  |  35.00  |  17,500.00
IQF Fish Fillet  |  030389  |  250 Cartons  |  29.40  |  7,350.00
"""


def test_title_line_is_not_read_as_a_field_label():
    f = extract_fields_from_text(_PO_EN, DocumentType.PURCHASE_ORDER)
    assert f["po_number"].value == "PO-2026-5001"
    assert f["grand_total"].value == "USD 24,850.00"
    assert f["currency"].value == "USD"


def test_wide_table_rows_become_line_items_not_fields():
    raw = RawDocument(name="ci.docx", pages=[RawPage(0, _INVOICE_EN)], text_is_native=True)
    ex = TextExtractor().extract(raw, DocumentType.COMMERCIAL_INVOICE)
    assert "description" not in ex.fields and "frozen_shrimp" not in ex.fields
    assert [li.fields["amount"].value for li in ex.line_items] == ["17,500.00", "7,350.00"]
    assert ex.fields["grand_total"].value == "24,850.00"


def test_air_waybill_origin_destination_are_airports():
    text = "House Air Waybill\nHAWB No.  |  HAWB-1\nOrigin  |  HYD\nDestination  |  DXB\n"
    f = extract_fields_from_text(text, DocumentType.HOUSE_AIR_WAYBILL)
    assert f["airport_of_departure"].value == "HYD"
    assert f["airport_of_destination"].value == "DXB"


def test_label_on_previous_line_and_dates_are_not_labels():
    text = "AWS Statement\nStatement Date:\nOctober 1 , 2026\nTOTAL AMOUNT DUE BY October 1 , 2026\n"
    f = extract_fields_from_text(text, DocumentType.UNKNOWN)
    assert f["statement_date"].value == "October 1 , 2026"
    assert "october" not in f and "total_amount_due_by_october" not in f


def test_reversed_arabic_pdf_text_is_detected():
    from app.ingest import _looks_reversed

    assert _looks_reversed("باسحلا فشك مقر\nيحاتتفالا ديصرلا\nيماتخلا ديصرلا")
    assert not _looks_reversed("رقم كشف الحساب\nالرصيد الافتتاحي\nالرصيد الختامي")


def test_arabic_titles_classify_to_the_right_type():
    from app.pipeline.classify import classify

    cases = {
        "شهادة ضريبة السلع والخدمات": DocumentType.GST_CERTIFICATE,
        "أمر شراء": DocumentType.PURCHASE_ORDER,
        "فاتورة المورد\nمرجع أمر الشراء": DocumentType.VENDOR_INVOICE,
        "إشعار استلام البضائع\nمرجع أمر الشراء": DocumentType.GOODS_RECEIPT_NOTE,
        "بوليصة شحن جوي فرعية": DocumentType.HOUSE_AIR_WAYBILL,
    }
    for text, expected in cases.items():
        assert classify(RawDocument(name="x", pages=[RawPage(0, text)])) == expected


def test_unreadable_uploads_never_crash_or_show_sample_data():
    import os

    from app.ingest import ingest_upload
    from app.pipeline.orchestrator import process_document

    for name, content in {
        "broken.jpg": os.urandom(3000),
        "empty.pdf": b"",
        "corrupt.docx": os.urandom(3000),
        "binary.bin": os.urandom(3000),
    }.items():
        sub = process_document(ingest_upload(name, content)).sub_documents[0]
        assert sub.engine != "mock" and sub.extraction.engine != "mock", name
        assert sub.route.value != "auto_approved", name


def test_auto_approve_policy_never_overrides_a_failed_critical_rule():
    from app.business_rules import apply_policy
    from app.models import RouteDecision

    decision, _ = apply_policy(DocumentType.HOUSE_AIR_WAYBILL, RouteDecision.NEEDS_REVIEW, 0.97, critical_failed=True)
    assert decision == RouteDecision.NEEDS_REVIEW
    decision, _ = apply_policy(DocumentType.HOUSE_AIR_WAYBILL, RouteDecision.NEEDS_REVIEW, 0.97)
    assert decision == RouteDecision.AUTO_APPROVED


def _docx_bytes(title: str, rows: list[tuple[str, str]]) -> bytes:
    import io

    import docx

    d = docx.Document()
    d.add_paragraph(title)
    table = d.add_table(rows=0, cols=2)
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = k, v
    d.add_paragraph("Authorised signatory")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_word_upload_gets_a_page_picture_and_traced_values():
    from app.ingest import ingest_upload
    from app.pipeline.orchestrator import process_document

    for title, rows in {
        "Certificate of Origin": [("Certificate No.", "COO-77"), ("Country of Origin", "India")],
        "شهادة منشأ": [("رقم الشهادة", "شهادة-منشأ-001"), ("بلد المنشأ", "الهند")],
    }.items():
        raw = ingest_upload("coo.docx", _docx_bytes(title, rows))
        assert raw.kind == "docx"
        assert raw.preview and raw.preview[:4] == b"\x89PNG"
        assert not raw.has_image  # display-only: extraction still reads the text
        # Document order: title, table rows, then the closing paragraph.
        assert raw.text.splitlines()[0] == title
        assert raw.text.splitlines()[-1] == "Authorised signatory"
        sub = process_document(raw).sub_documents[0]
        assert sub.extraction.engine == "text:docx/text"
        values = [f for f in sub.extraction.fields.values() if f.value in {v for _, v in rows}]
        assert values and all(f.bbox and 0 <= f.bbox.x < 1 and 0 <= f.bbox.y < 1 for f in values)


def test_pdf_values_are_traced_to_their_place_on_the_page():
    import pymupdf

    from app.ingest import ingest_upload
    from app.pipeline.orchestrator import process_document

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Bank Statement", fontsize=14)
    page.insert_text((72, 120), "Statement No.: BS-2026-001", fontsize=11)
    page.insert_text((72, 300), "Closing Balance: 142,450.00", fontsize=11)
    raw = ingest_upload("s.pdf", doc.tobytes())
    sub = process_document(raw).sub_documents[0]
    f = sub.extraction.fields
    assert f["statement_no"].bbox and f["closing_balance"].bbox
    # Lower on the page = larger y.
    assert f["closing_balance"].bbox.y > f["statement_no"].bbox.y


def test_preview_endpoint_helper_handles_every_kind():
    import os

    import pymupdf

    from app.ingest import preview_png

    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "Hello")
    assert preview_png("a.pdf", pdf.tobytes())[1] == "image/png"
    assert preview_png("a.docx", _docx_bytes("Invoice", [("Invoice No.", "1")]))[1] == "image/png"
    png = preview_png("a.pdf", pdf.tobytes())[0]
    assert preview_png("a.png", png, "image/png") == (png, "image/png")
    assert preview_png("x.pdf", os.urandom(2000)) is None


def test_trade_analytics_merge_languages_and_ignore_non_countries():
    from app.analytics import real_geo

    def doc(fields: dict) -> dict:
        return {"sub_documents": [{"min_confidence": 0.9, "extraction": {"fields": {
            k: {"value": v} for k, v in fields.items()}}}]}

    geo = real_geo([
        doc({"country_of_origin": "India", "destination_country": "United Arab Emirates"}),
        doc({"country_of_origin": "الهند", "destination_country": "الإمارات العربية المتحدة"}),
        doc({"consignee": "Gulf Foods Import Co.", "airport_of_destination": "DXB"}),
    ])
    assert geo["by_origin"] == [{"country": "India", "code": "IN", "count": 2}]
    assert geo["destinations"] == [{"country": "United Arab Emirates", "count": 2}]
    lane = geo["trade_lanes"][0]
    assert (lane["origin_code"], lane["destination_code"], lane["count"]) == ("IN", "AE", 2)
